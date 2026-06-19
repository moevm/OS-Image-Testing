from __future__ import annotations

import json
import logging
import shlex
from collections.abc import Callable, Iterable, Mapping, Sequence
from datetime import date, datetime
from pathlib import Path
from typing import TYPE_CHECKING, Any, Final, TypedDict

from openpyxl import Workbook
from sqlalchemy import create_engine, inspect, select
from sqlalchemy.orm import Session

from imgtests.constant import DISTRIBUTION_DESCRIPTIONS
from imgtests.database.models.configuration import ConfigurationBase
from imgtests.database.models.experiment import ExperimentBase
from imgtests.database.models.util_run_result import UtilRunResult
from imgtests.reporting.cli import (
    DISTRO_COMPARISON_COMMAND,
    DISTRO_COMPARISON_STATUS_COMMAND,
    EXPORT_TABLES,
    parse_args,
)
from imgtests.reporting.distro_comparison_export import (
    DistroComparisonExportOptions,
    export_distro_comparison_to_excel,
)
from imgtests.reporting.distro_comparison_status_export import (
    DistroComparisonStatusOptions,
    export_distro_comparison_status_to_excel,
)
from imgtests.types import Distro

if TYPE_CHECKING:
    import argparse

    from sqlalchemy.engine import Engine

logger = logging.getLogger(__name__)

CONFIGURATION_SHEET: Final = "configuration"
DEFAULT_OUTPUT_FILENAME: Final = "report.xlsx"

JSON_COLUMNS: Final = {
    "util_run_result": {"result"},
}

EXCEL_MAX_COLUMNS: Final = 16_384
EXCEL_MAX_ROWS: Final = 1_048_576

UNNAMED_UTIL_RUN_RESULT_SHEET: Final = "unnamed_util_run_result"

NOISY_COLUMN_PARTS: Final = frozenset(
    {
        "result",
        "metrics",
        "test_type",
        "detailed",
        "start",
        "end",
        "connected",
        "summary",
        "time",
    },
)

COLUMN_PART_REPLACEMENTS: Final = {
    "bits_per_second": "bps",
    "bogo_ops_per_sec_real_time": "bogo_ops_per_sec_real",
    "bogo_ops_per_sec_usr_sys_time": "bogo_ops_per_sec_cpu",
    "cpu_utilization_percent": "cpu",
    "sum_sent": "sent",
    "sum_received": "received",
    "test_start": "test",
    "num_streams": "streams",
    "numjobs": "num_jobs",
    "blksize": "block_size",
    "bs": "block_size",
    "ioengine": "io_engine",
    "iodepth": "io_depth",
    "rw": "read_write",
    "duration": "duration_sec",
    "usr_cpu": "user_cpu_pct",
    "sys_cpu": "system_cpu_pct",
    "majf": "major_faults",
    "minf": "minor_faults",
    "ctx": "context_switches",
}


def export_database_to_excel(
    engine: Engine,
    output_path: Path,
    configuration_ids: dict[str, int],
    tables: Sequence[str] = tuple(EXPORT_TABLES),
) -> Path:
    output_path = prepare_output_path(output_path)
    distributions = build_distribution_records(configuration_ids)
    unsupported_tables = [table for table in tables if table not in EXPORT_TABLES]

    if unsupported_tables:
        joined_tables = ", ".join(unsupported_tables)
        msg = f"Unsupported export tables: {joined_tables}"
        raise ValueError(msg)

    existing_tables = set(inspect(engine).get_table_names())
    required_tables = list(tables)

    if "util_run_result" in tables and "experiment" not in required_tables:
        required_tables.append("experiment")

    if any(table in tables for table in ("experiment", "util_run_result")):
        required_tables.append("configuration")

    missing_tables = [table for table in required_tables if table not in existing_tables]

    if missing_tables:
        joined_tables = ", ".join(missing_tables)
        msg = f"Tables not found in database: {joined_tables}"
        raise RuntimeError(msg)

    worksheets: list[WorksheetData] = []

    if any(table in tables for table in ("experiment", "util_run_result")):
        worksheets.append(
            flatten_table_records(
                CONFIGURATION_SHEET,
                fetch_configuration_records(engine, distributions),
            ),
        )

    for table in tables:
        records = fetch_table_records(engine, table, distributions)

        if table == "util_run_result":
            worksheets.extend(flatten_util_run_results(records))
        else:
            worksheets.append(flatten_table_records(table, records))

    write_xlsx(output_path, worksheets)
    return output_path


def fetch_table_records(
    engine: Engine,
    table: str,
    distributions: Mapping[str, DistributionRecord],
) -> list[dict[str, Any]]:
    if table == "experiment":
        return fetch_experiment_records(engine, distributions)

    if table == "util_run_result":
        return fetch_util_run_result_records(engine, distributions)

    msg = f"Unsupported export table: {table}"
    raise ValueError(msg)


def fetch_experiment_records(
    engine: Engine,
    distributions: Mapping[str, DistributionRecord],
) -> list[dict[str, Any]]:
    query = (
        select(
            ConfigurationBase.os.label("configuration_os"),
            ExperimentBase.experiment_id,
            ExperimentBase.config_id,
            ExperimentBase.description,
            ExperimentBase.type,
            ExperimentBase.started_at,
            ExperimentBase.ended_at,
            ExperimentBase.tests_total,
            ExperimentBase.tests_passed,
            ExperimentBase.tests_failed,
            ExperimentBase.tests_broken,
            ExperimentBase.tests_skipped,
        )
        .select_from(ExperimentBase)
        .outerjoin(ExperimentBase.configuration)
        .order_by(ExperimentBase.experiment_id)
    )

    with Session(engine) as session:
        result = session.execute(query)
        return [add_configuration_id(dict(row), distributions) for row in result.mappings()]


def fetch_util_run_result_records(
    engine: Engine,
    distributions: Mapping[str, DistributionRecord],
) -> list[dict[str, Any]]:
    query = (
        select(
            ConfigurationBase.os.label("configuration_os"),
            UtilRunResult.id,
            UtilRunResult.experiment_id,
            UtilRunResult.util_type,
            UtilRunResult.command,
            UtilRunResult.result,
            UtilRunResult.description,
            UtilRunResult.started_at,
            UtilRunResult.ended_at,
        )
        .select_from(UtilRunResult)
        .outerjoin(UtilRunResult.experiment)
        .outerjoin(ExperimentBase.configuration)
        .order_by(UtilRunResult.id)
    )

    with Session(engine) as session:
        result = session.execute(query)
        return [add_configuration_id(dict(row), distributions) for row in result.mappings()]


def fetch_configuration_records(
    _engine: Engine,
    distributions: Mapping[str, DistributionRecord],
) -> list[dict[str, Any]]:
    return [dict(record) for record in distributions.values()]


def add_configuration_id(
    record: dict[str, Any],
    distributions: Mapping[str, DistributionRecord],
) -> dict[str, Any]:
    os_name = record.pop("configuration_os", None)
    return {"configuration_id": configuration_id_from_os(os_name, distributions), **record}


def configuration_id_from_os(
    os_name: Any,
    distributions: Mapping[str, DistributionRecord],
) -> int | str:
    distribution_name = distribution_name_from_os(os_name)

    if not distribution_name:
        return ""

    return int(distributions[distribution_name]["configuration_id"])


class DistributionRecord(TypedDict):
    configuration_id: int
    distribution_description: str


def build_distribution_records(
    configuration_ids: dict[str, int],
) -> dict[str, DistributionRecord]:
    ids = {
        distribution_name.strip().lower(): configuration_id
        for distribution_name, configuration_id in configuration_ids.items()
    }
    unsupported_distributions = [
        distribution_name
        for distribution_name in ids
        if distribution_name not in DISTRIBUTION_DESCRIPTIONS
    ]

    if unsupported_distributions:
        valid_names = ", ".join(DISTRIBUTION_DESCRIPTIONS)
        joined_distributions = ", ".join(unsupported_distributions)
        msg = f"Unsupported distributions: {joined_distributions}. Available: {valid_names}."
        raise ValueError(msg)

    missing_distributions = [
        distribution_name
        for distribution_name in DISTRIBUTION_DESCRIPTIONS
        if distribution_name not in ids
    ]

    if missing_distributions:
        joined_distributions = ", ".join(missing_distributions)
        msg = f"Missing configuration ids for distributions: {joined_distributions}."
        raise ValueError(msg)

    return {
        distribution_name: DistributionRecord(
            configuration_id=ids[distribution_name],
            distribution_description=distribution_description,
        )
        for distribution_name, distribution_description in DISTRIBUTION_DESCRIPTIONS.items()
    }


def distribution_name_from_os(os_name: Any) -> str:
    value = str(os_name or "").strip().lower()

    if Distro.POKY.value in value or "yocto" in value:
        return "poky"
    if "suse" in value:
        return "suse"

    return ""


def flatten_table_records(
    table: str,
    records: Iterable[Mapping[str, Any]],
    json_columns: set[str] | None = None,
) -> WorksheetData:
    headers: dict[str, None] = {}
    rows: list[dict[str, Any]] = []
    json_columns = json_columns if json_columns is not None else JSON_COLUMNS.get(table, set())

    for record in records:
        row = flatten_record(record, json_columns)

        for header, value in row.items():
            if not is_empty_cell(value):
                headers.setdefault(header)

        rows.append(row)

    return WorksheetData(name=table, headers=list(headers), rows=rows)


def flatten_record(record: Mapping[str, Any], json_columns: set[str]) -> dict[str, Any]:
    return flatten_record_with_column_names(record, json_columns)


def flatten_record_with_column_names(
    record: Mapping[str, Any],
    json_columns: set[str],
    column_name: Callable[[str], str] | None = None,
) -> dict[str, Any]:
    row: dict[str, Any] = {}

    for column, value in record.items():
        if column in json_columns:
            for key, nested_value in flatten_json(column, decode_json(value)):
                add_flat_cell(row, resolve_column_name(key, column_name), nested_value)
        elif column == "command":
            row[column] = value
        else:
            add_flat_cell(row, resolve_column_name(column, column_name), value)

    return row


def flatten_util_run_results(
    records: Iterable[Mapping[str, Any]],
) -> list[WorksheetData]:
    prepared_rows: list[dict[str, Any]] = []

    for record in records:
        row = flatten_util_run_result(record)
        headers = row_non_empty_headers(row)

        prepared_rows.append(
            {
                "row": row,
                "headers": headers,
                "test_name": str(row.get("test_name") or ""),
            },
        )

    if not prepared_rows:
        return [flatten_table_records("util_run_result", [])]

    grouped_records: dict[tuple[str, tuple[str, ...]], list[dict[str, Any]]] = {}

    for prepared_row in prepared_rows:
        sheet_name = util_run_result_sheet_name(prepared_row)
        headers = prepared_row["headers"]
        row = prepared_row["row"]

        group_key = (sheet_name, headers)
        grouped_records.setdefault(group_key, []).append(row)

    worksheets: list[WorksheetData] = []

    for (sheet_name, headers), rows in grouped_records.items():
        worksheets.extend(split_rows_to_worksheets(sheet_name, list(headers), rows))

    return worksheets


def flatten_util_run_result(record: Mapping[str, Any]) -> dict[str, Any]:
    result = decode_json(record.get("result"))
    test_name = util_run_result_test_name(record, result)
    tool_name = detect_utility(result)
    record_without_configuration = dict(record)

    row: dict[str, str] = {
        "configuration_id": record_without_configuration.pop("configuration_id", ""),
        "test_name": test_name,
        "tool_name": tool_name,
    }
    row.update(
        flatten_record_with_column_names(
            record_without_configuration,
            JSON_COLUMNS["util_run_result"],
            readable_util_column_name,
        ),
    )

    return row


def util_run_result_sheet_name(prepared_row: Mapping[str, Any]) -> str:
    sheet_name = compact_sheet_part(str(prepared_row.get("test_name") or ""))

    if sheet_name:
        return sheet_name

    return UNNAMED_UTIL_RUN_RESULT_SHEET


def split_rows_to_worksheets(
    sheet_name: str,
    headers: list[str],
    rows: list[dict[str, Any]],
) -> list[WorksheetData]:
    max_data_rows = EXCEL_MAX_ROWS - 1

    if len(rows) <= max_data_rows:
        return [WorksheetData(sheet_name, headers, rows)]

    worksheets: list[WorksheetData] = []

    for index, start in enumerate(range(0, len(rows), max_data_rows), start=1):
        chunk = rows[start : start + max_data_rows]
        worksheets.append(WorksheetData(f"{sheet_name}_part_{index}", headers, chunk))

    return worksheets


def resolve_column_name(key: str, column_name: Callable[[str], str] | None) -> str:
    if column_name is None:
        return key

    return column_name(key)


def row_non_empty_headers(row: Mapping[str, Any]) -> tuple[str, ...]:
    return tuple(header for header, value in row.items() if not is_empty_cell(value))


def add_flat_cell(row: dict[str, Any], key: str, value: Any) -> None:
    if key in row:
        if row[key] == value:
            return

        key = next_duplicate_key(row, key)

    if is_empty_cell(value):
        row[key] = value
        return

    if isinstance(value, str):
        add_text_cells(row, key, value)
        return

    row[key] = value


def next_duplicate_key(row: Mapping[str, Any], key: str) -> str:
    index = 2
    duplicate_key = f"{key}_{index}"

    while duplicate_key in row:
        index += 1
        duplicate_key = f"{key}_{index}"

    return duplicate_key


def add_text_cells(row: dict[str, Any], key: str, value: str) -> None:
    parts = split_long_text(value)

    if len(parts) == 1:
        row[key] = parts[0]
        return

    for index, part in enumerate(parts, start=1):
        row[f"{key}_part_{index}"] = part


def split_long_text(value: str) -> list[str]:
    lines = [line.strip() for line in value.splitlines() if line.strip()]

    if len(lines) > 1:
        return lines

    return [value]


def util_run_result_group(record: Mapping[str, Any]) -> str:
    return util_run_result_test_name(record, decode_json(record.get("result")))


def util_run_result_test_name(record: Mapping[str, Any], result: Any) -> str:
    description = str(record.get("description") or "").strip()

    if description:
        return description

    result_name = result_based_test_name(result)

    if result_name:
        return result_name

    return command_name_from_record(record)


def result_based_test_name(result: Any) -> str:
    if not isinstance(result, dict):
        return ""

    tool = str(result.get("tool", "")).strip()
    if tool:
        return result_tool_test_name(tool, result.get("test_type"))

    return ""


def result_tool_test_name(tool: str, test_type: Any) -> str:
    if not isinstance(test_type, dict):
        return tool

    for field in ("name", "stressor", "protocol"):
        value = str(test_type.get(field, "")).strip()
        if value and value != "unknown":
            return f"{tool}_{value}"

    return tool


def detect_utility(result: Any) -> str:
    if isinstance(result, dict):
        return str(result.get("tool", "")).strip()

    return ""


def command_name_from_record(record: Mapping[str, Any]) -> str:
    command = str(record.get("command", "")).strip()

    if not command:
        return ""

    return command_name(command)


def command_name(command: str) -> str:
    try:
        parts = shlex.split(command)
    except ValueError:
        parts = command.split()

    for part in parts:
        executable = Path(part).name

        if (
            executable
            and executable not in {"sudo", "env"}
            and not executable.startswith("-")
            and "=" not in executable
        ):
            return executable

    return ""


def readable_util_column_name(key: str) -> str:
    return normalize_metric_column(key)


def normalize_metric_column(path: str) -> str:
    parts: list[str] = []
    for raw_part in column_path_parts(path):
        if raw_part.isdecimal():
            continue
        part = COLUMN_PART_REPLACEMENTS.get(raw_part, raw_part)
        if part in NOISY_COLUMN_PARTS:
            continue
        if parts and parts[-1] == part:
            continue
        parts.append(part)

    return "_".join(parts) or "value"


def column_path_parts(path: str) -> list[str]:
    parts: list[str] = []
    current: list[str] = []

    for char in path:
        if char.isalnum() or char == "_":
            current.append(char.lower())
            continue
        if current:
            parts.append("".join(current))
            current.clear()

    if current:
        parts.append("".join(current))

    return parts


def compact_sheet_part(value: str) -> str:
    return "_".join(column_path_parts(value))


def decode_json(value: Any) -> Any:
    if isinstance(value, str):
        try:
            return json.loads(value)
        except json.JSONDecodeError:
            return value

    return value


def flatten_json(prefix: str, value: Any) -> list[tuple[str, Any]]:
    decoded_value = decode_nested_json(value)

    if decoded_value is not value:
        return flatten_json(prefix, decoded_value)

    if isinstance(value, dict):
        if not value:
            return [(prefix, None)]

        flattened: list[tuple[str, Any]] = []
        for key, nested_value in value.items():
            nested_prefix = f"{prefix}.{key}"
            flattened.extend(flatten_json(nested_prefix, nested_value))

        return flattened

    if isinstance(value, list):
        if not value:
            return [(prefix, None)]

        flattened = []

        for index, nested_value in enumerate(value):
            nested_prefix = f"{prefix}[{index}]"
            flattened.extend(flatten_json(nested_prefix, nested_value))

        return flattened

    return [(prefix, value)]


def decode_nested_json(value: Any) -> Any:
    if not isinstance(value, str):
        return value

    stripped_value = value.strip()

    if not stripped_value or stripped_value[0] not in "[{":
        return value

    try:
        return json.loads(stripped_value)
    except json.JSONDecodeError:
        return value


def is_empty_cell(value: Any) -> bool:
    return value is None or value in ("", {}, [])


class WorksheetData:
    def __init__(self, name: str, headers: list[str], rows: list[dict[str, Any]]) -> None:
        if len(headers) > EXCEL_MAX_COLUMNS:
            msg = (
                f"Worksheet '{name}' has {len(headers)} columns, "
                f"Excel supports {EXCEL_MAX_COLUMNS}."
            )
            raise ValueError(msg)

        if len(rows) + 1 > EXCEL_MAX_ROWS:
            msg = f"Worksheet '{name}' has {len(rows) + 1} rows, Excel supports {EXCEL_MAX_ROWS}."
            raise ValueError(msg)

        self.name = name
        self.headers = headers
        self.rows = rows


def write_xlsx(output_path: Path, worksheets: Sequence[WorksheetData]) -> None:
    output_path = prepare_output_path(output_path)
    sheet_names = unique_sheet_names(worksheet.name for worksheet in worksheets)

    workbook = Workbook()
    workbook.remove(workbook.active)

    for sheet_name, worksheet in zip(sheet_names, worksheets, strict=True):
        sheet = workbook.create_sheet(title=sheet_name)
        sheet.append(worksheet.headers)

        for row in worksheet.rows:
            sheet.append([value_to_cell(row.get(header)) for header in worksheet.headers])

    workbook.save(output_path)


def prepare_output_path(output_path: Path) -> Path:
    output_path = output_path.expanduser()

    if (output_path.exists() and output_path.is_dir()) or not output_path.suffix:
        output_path /= DEFAULT_OUTPUT_FILENAME
    elif output_path.suffix.lower() != ".xlsx":
        msg = f"Output file must have .xlsx extension: {output_path}"
        raise ValueError(msg)

    output_parent = output_path.parent

    if output_parent.exists() and not output_parent.is_dir():
        msg = f"Output parent path is not a directory: {output_parent}"
        raise ValueError(msg)

    try:
        output_parent.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        msg = f"Cannot create output directory '{output_parent}': {exc}"
        raise ValueError(msg) from exc

    if output_path.exists() and output_path.is_dir():
        msg = f"Output path is a directory: {output_path}"
        raise ValueError(msg)

    if output_path.exists() and not output_path.is_file():
        msg = f"Output path is not a regular file: {output_path}"
        raise ValueError(msg)

    return output_path


def unique_sheet_names(names: Iterable[str]) -> list[str]:
    used_names: set[str] = set()
    result: list[str] = []

    for name in names:
        sheet_name = sanitize_sheet_name(name)
        original_name = sheet_name
        suffix = 1

        while sheet_name in used_names:
            suffix_text = f"_{suffix}"
            sheet_name = f"{original_name[: 31 - len(suffix_text)]}{suffix_text}"
            suffix += 1

        used_names.add(sheet_name)
        result.append(sheet_name)

    return result


def sanitize_sheet_name(name: str) -> str:
    cleaned_name = "".join("_" if char in r":\/?*[]" else char for char in name).strip("'")

    return (cleaned_name or "Sheet")[:31]


def value_to_cell(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()

    if isinstance(value, date):
        return value.isoformat()

    if isinstance(value, Mapping | list):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)

    return value


def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    args = parse_args()

    if args.command == DISTRO_COMPARISON_COMMAND:
        export_distro_comparison_command(args)
        return

    if args.command == DISTRO_COMPARISON_STATUS_COMMAND:
        export_distro_comparison_status_command(args)
        return

    engine = create_engine(args.db_url)

    try:
        output_path = export_database_to_excel(
            engine=engine,
            output_path=args.output,
            tables=args.tables,
            configuration_ids=dict(args.configuration_ids),
        )
    except ValueError as exc:
        raise SystemExit(str(exc)) from exc

    logger.info("Exported tables %s to %s", ", ".join(args.tables), output_path)


def export_distro_comparison_command(args: argparse.Namespace) -> None:
    output_path = export_distro_comparison_to_excel(
        input_path=args.input,
        options=DistroComparisonExportOptions(
            output_path=args.output,
            experiment_ids=args.experiment_ids,
            max_charts=args.max_charts,
            charts_per_sheet=args.charts_per_sheet,
            copy_source_sheets=args.copy_source_sheets,
            include_comparison=args.include_comparison,
        ),
    )
    logger.info("Exported comparison XLSX file: %s", output_path)


def export_distro_comparison_status_command(args: argparse.Namespace) -> None:
    output_path = export_distro_comparison_status_to_excel(
        input_path=args.input,
        options=DistroComparisonStatusOptions(
            output_path=args.output,
            experiment_ids=args.experiment_ids,
            epsilon_percent=args.epsilon_percent,
        ),
    )
    logger.info("Exported comparison status XLSX file: %s", output_path)


if __name__ == "__main__":
    main()
